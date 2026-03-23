"""
app/reports/expense_transmittal_docx.py

Generate "ΔΙΑΒΙΒΑΣΤΙΚΟ ΔΑΠΑΝΗΣ" as DOCX bytes using a Word template
and placeholder replacement.

SOURCE OF TRUTH
---------------
This implementation is aligned strictly to the provided current template and
the current ServiceUnit / Procurement contract.

IMPORTANT FIELD MAPPING RULES
-----------------------------
The uploaded DOCX contains placeholders that must be resolved from the current
domain objects only.

Current source-of-truth fields used:
- procurement.hop_approval
- procurement.hop_preapproval
- procurement.aay
- procurement.protocol_number
- procurement.committee.identity_text
- procurement.invoice_number
- procurement.invoice_date
- procurement.invoice_receipt_date
- procurement.identity_prosklisis
- service_unit.description
- service_unit.phone
- service_unit.region
- service_unit.curator
- service_unit.application_admin_directory

IMPORTANT RENDERING RULES
-------------------------
1. {{SUPPORTING_DOCUMENTS_BLOCK}} must be placed alone in its own paragraph.
2. The supporting-documents block is rendered as separate paragraphs.
3. There must be NO blank paragraphs between η., θ., ι., ...
4. Each generated line must begin with exactly two tab characters.
5. Paragraph tab stops / paragraph properties are cloned from paragraph 'ζ.'.
6. Placeholder replacement must work even when placeholders are split across runs.
7. Run-level formatting in the template must be preserved as much as possible.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

from docx import Document

from .common.amounts import money_words_el
from .common.docx_utils import (
    clear_paragraph_runs,
    clone_paragraph_properties,
    copy_run_style,
    insert_paragraph_after,
    iter_document_paragraphs,
    replace_placeholders_everywhere,
    set_global_font_arial_12,
)
from .common.domain import (
    resolve_analysis_total,
    resolve_document_total,
    resolve_document_total_value,
    resolve_is_services,
    winner_supplier_line,
)
from .common.formatting import (
    format_date_ddmmyyyy,
    safe_text,
    short_date_el,
    upper_service_name,
)


def _template_path() -> Path:
    """
    Resolve the DOCX template path for the expense transmittal document.
    """
    app_dir = Path(__file__).resolve().parents[1]
    return app_dir / "templates" / "docx" / "expense_transmittal_template.docx"


def _resolve_proc_type(procurement: Any) -> str:
    """
    Resolve whether the procurement concerns services or goods.
    """
    return "παροχής υπηρεσιών" if resolve_is_services(procurement) else "προμήθειας υλικών"


def _resolve_committee_description(procurement: Any) -> str:
    """
    Resolve the linked committee identity text.

    The current document must show committee identity text.
    """
    committee = getattr(procurement, "committee", None)
    if committee is None:
        return "—"
    return safe_text(getattr(committee, "identity_text", None))


def _resolve_application_admin(service_unit: Any) -> str:
    """
    Resolve application administrator text from current ServiceUnit model.
    """
    return safe_text(getattr(service_unit, "curator", None))


def _resolve_application_admin_directory(service_unit: Any) -> str:
    """
    Resolve free-text application-admin directory from current ServiceUnit model.
    """
    return safe_text(getattr(service_unit, "application_admin_directory", None))


def _greek_enumeration_labels() -> list[str]:
    """
    Ordered Greek labels for the supporting-documents block.
    """
    return [
        "η.",
        "θ.",
        "ι.",
        "ια.",
        "ιβ.",
        "ιγ.",
        "ιδ.",
        "ιε.",
        "ιστ.",
        "ιζ.",
        "ιη.",
        "ιθ.",
        "κ.",
    ]


def _build_supporting_document_items(procurement: Any, analysis: dict[str, Any]) -> list[str]:
    """
    Build supporting-document lines according to amount thresholds.
    """
    total_amount = resolve_analysis_total(procurement, analysis)

    full_items = [
        f"Πρόσκληση Υποβολής Προσφοράς με {safe_text(getattr(procurement, 'identity_prosklisis', None))}.",
        "Βεβαίωση ΙΒΑΝ.",
        "Υπεύθυνη δήλωση στοιχείων επικοινωνίας και τραπεζικών στοιχείων.",
        "Πιστοποιητικό Εκπροσώπησης.",
        "Υπεύθυνη Δήλωση μη υποχρέωσης ένταξης στο Εθνικό Μητρώο Παραγωγών.",
        "Υπεύθυνη Δήλωση μη δωροδοκίας.",
        "Αντίγραφο Ποινικού Μητρώου.",
        "Αποδεικτικό Φορολογικής Ενημερότητας.",
        "Αποδεικτικό Ασφαλιστικής Ενημερότητας.",
    ]

    if total_amount < Decimal("1500"):
        return [
            "Βεβαίωση ΙΒΑΝ.",
            "Υπεύθυνη δήλωση στοιχείων επικοινωνίας και τραπεζικών στοιχείων.",
        ]

    if total_amount < Decimal("2500"):
        return [
            "Βεβαίωση ΙΒΑΝ.",
            "Υπεύθυνη δήλωση στοιχείων επικοινωνίας και τραπεζικών στοιχείων.",
            "Αποδεικτικό Φορολογικής Ενημερότητας.",
        ]

    return full_items


def _find_reference_paragraph_for_supporting_block(doc: Document):
    """
    Find the paragraph that starts with 'ζ.' and use it as the formatting source.
    """
    for paragraph in iter_document_paragraphs(doc):
        text = (paragraph.text or "").strip()
        if text.startswith("ζ."):
            return paragraph

    return None


def _render_supporting_documents_block_in_paragraph(
    paragraph,
    items: list[str],
    formatting_source_paragraph=None,
) -> None:
    """
    Render supporting documents as consecutive paragraphs while cloning
    paragraph formatting from the 'ζ.' paragraph.
    """
    labels = _greek_enumeration_labels()
    source_paragraph = formatting_source_paragraph or paragraph
    template_run = source_paragraph.runs[0] if source_paragraph.runs else (
        paragraph.runs[0] if paragraph.runs else None
    )

    clear_paragraph_runs(paragraph)
    clone_paragraph_properties(source_paragraph, paragraph)

    anchor = paragraph

    for idx, item in enumerate(items):
        if idx >= len(labels):
            raise ValueError("Not enough Greek enumeration labels for supporting documents block.")

        target = anchor if idx == 0 else insert_paragraph_after(anchor)
        clone_paragraph_properties(source_paragraph, target)

        run = target.add_run(f"\t\t{labels[idx]}\t{item}")
        copy_run_style(template_run, run)

        anchor = target


def _render_supporting_documents_block(doc: Document, placeholder: str, items: list[str]) -> None:
    """
    Replace the placeholder paragraph with the rendered supporting-documents block.
    """
    formatting_source_paragraph = _find_reference_paragraph_for_supporting_block(doc)

    for paragraph in iter_document_paragraphs(doc):
        if placeholder in paragraph.text:
            _render_supporting_documents_block_in_paragraph(
                paragraph=paragraph,
                items=items,
                formatting_source_paragraph=formatting_source_paragraph,
            )
            return


@dataclass(frozen=True)
class ExpenseTransmittalConstants:
    """
    Future-proof constants container for expense transmittal generation.
    """
    pass


def build_expense_transmittal_docx(
    *,
    procurement: Any,
    service_unit: Any,
    winner: Optional[Any],
    analysis: dict[str, Any],
    constants: Optional[ExpenseTransmittalConstants] = None,
) -> bytes:
    """
    Build the expense transmittal DOCX and return it as bytes.
    """
    _ = constants

    template_path = _template_path()
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    doc = Document(str(template_path))

    supporting_items = _build_supporting_document_items(procurement, analysis)

    mapping: dict[str, str] = {
        "{{SERVICE_UNIT_NAME}}": upper_service_name(
            safe_text(getattr(service_unit, "description", None), default="—")
        ),
        "{{APPLICATION_ADMIN_DIRECTORY}}": _resolve_application_admin_directory(service_unit),
        "{{APPLICATION_ADMIN}}": _resolve_application_admin(service_unit),
        "{{SERVICE_UNIT_PHONE}}": safe_text(getattr(service_unit, "phone", None)),
        "{{SERVICE_UNIT_REGION}}": safe_text(getattr(service_unit, "region", None)),
        "{{SHORT_DATE}}": short_date_el(),
        "{{PROC_TYPE}}": _resolve_proc_type(procurement),
        "{{procurement.hop_approval}}": safe_text(getattr(procurement, "hop_approval", None)),
        "{{procurement.hop_preapproval}}": safe_text(
            getattr(procurement, "hop_preapproval", None)
        ),
        "{{procurement. hop_preapproval}}": safe_text(
            getattr(procurement, "hop_preapproval", None)
        ),
        "{{procurement.aay}}": safe_text(getattr(procurement, "aay", None)),
        "{{procurement.protocol_number}}": safe_text(
            getattr(procurement, "protocol_number", None)
        ),
        "{{procurement. protocol_number}}": safe_text(
            getattr(procurement, "protocol_number", None)
        ),
        "{{procurement.committee_description}}": _resolve_committee_description(procurement),
        "{{WINNER_SUPPLIER_LINE}}": winner_supplier_line(winner),
        "{{procurement.invoice_number}}": safe_text(getattr(procurement, "invoice_number", None)),
        "{{procurement.invoice_date}}": format_date_ddmmyyyy(
            getattr(procurement, "invoice_date", None)
        ),
        "{{ML_TOTAL}}": resolve_document_total(procurement, analysis),
        "{{ML_TOTAL_WORDS}}": money_words_el(resolve_document_total_value(procurement, analysis)),
        "{{procurement.invoice_receipt_date}}": format_date_ddmmyyyy(
            getattr(procurement, "invoice_receipt_date", None)
        ),
        "{{procurement.identity_prosklisis}}": safe_text(
            getattr(procurement, "identity_prosklisis", None)
        ),
        "{{ProcurementCommittee.description}}": _resolve_committee_description(procurement),
        "{{procurement.invoice}}": safe_text(getattr(procurement, "invoice_number", None)),
        "{{procurement.date}}": format_date_ddmmyyyy(getattr(procurement, "invoice_date", None)),
        "{{MANAGER_SERVICE}}": _resolve_application_admin(service_unit),
    }

    replace_placeholders_everywhere(doc, mapping)

    _render_supporting_documents_block(
        doc=doc,
        placeholder="{{SUPPORTING_DOCUMENTS_BLOCK}}",
        items=supporting_items,
    )

    set_global_font_arial_12(doc)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def build_expense_transmittal_filename(
    *,
    procurement: Any,
    winner: Optional[Any],
) -> str:
    """
    Build a human-readable filename for the generated expense transmittal document.
    """
    supplier_name = safe_text(getattr(winner, "name", None), default="—")
    total_str = resolve_document_total(procurement, {})

    return f"Διαβιβαστικό Δαπάνης {supplier_name} {total_str}.docx"